#!/usr/bin/env python3
"""Validate the generated DOCX against the report-to-google-doc manifest."""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document


def main() -> int:
    if len(sys.argv) != 4:
        raise SystemExit("usage: validate_report_docx.py MANIFEST DOCX OUT_JSON")

    manifest_path, docx_path, out_path = map(Path, sys.argv[1:])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    document = Document(docx_path)

    paragraph_text = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    heading_text = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip() and p.style and p.style.name.startswith("Heading")
    ]
    expected_headings = [item["text"] for item in manifest.get("headings", [])]

    with zipfile.ZipFile(docx_path) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

    expected_tables = int(manifest.get("counts", {}).get("tables", 0))
    expected_images = int(manifest.get("counts", {}).get("images", 0))
    image_relationships = relationships.count("/image")
    missing_headings = [heading for heading in expected_headings if heading not in heading_text]

    checks = {
        "file_non_empty": docx_path.stat().st_size > 0,
        "title_present": manifest.get("title") in paragraph_text,
        "all_headings_present": not missing_headings,
        "heading_count_matches": len(heading_text) == len(expected_headings),
        "table_count_matches": len(document.tables) == expected_tables,
        "image_relationship_count_matches": image_relationships == expected_images,
        "no_unresolved_placeholders": not any("[[" in text or "]]" in text for text in paragraph_text),
    }
    result = {
        "status": "passed" if all(checks.values()) else "failed",
        "checks": checks,
        "docx_bytes": docx_path.stat().st_size,
        "paragraphs": len(document.paragraphs),
        "headings_expected": len(expected_headings),
        "headings_observed": len(heading_text),
        "missing_headings": missing_headings,
        "tables_expected": expected_tables,
        "tables_observed": len(document.tables),
        "image_relationships_expected": expected_images,
        "image_relationships_observed": image_relationships,
    }
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
